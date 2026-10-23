"""共通：総務省「人口推計」長期時系列（年齢各歳、1920〜2024年）を読んで、年×年齢×男女の表にする"""
from pathlib import Path
import io, re, logging, warnings
import numpy as np, pandas as pd
import matplotlib.pyplot as plt
from matplotlib import font_manager
warnings.filterwarnings("ignore")
logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
for f in font_manager.findSystemFonts():
    if "NotoSansCJK" in f or "NotoSans-CJK" in f:
        font_manager.fontManager.addfont(f)
plt.rcParams.update({"font.family": ["Noto Sans CJK JP", "IPAexGothic", "Yu Gothic", "Meiryo", "sans-serif"],
                     "axes.spines.top": False, "axes.spines.right": False, "font.size": 11, "axes.titlesize": 13,
                     "figure.dpi": 150, "savefig.bbox": "tight"})
ROOT = Path(__file__).resolve().parents[1]
RAW, DATA, FIG, RES = ROOT / "data" / "raw", ROOT / "data", ROOT / "figures", ROOT / "results"
FIG.mkdir(exist_ok=True); RES.mkdir(exist_ok=True)
BLUE, ORANGE, GRAY, INK, LIGHT, GREEN, PURPLE = "#2a78d6", "#eb6834", "#a3a29c", "#0b0b0b", "#a9c9ef", "#3aa07a", "#7a5bd6"
SOURCE = "出典：総務省統計局「人口推計」長期時系列データ・各年10月1日現在人口（e-Stat）をもとに作成者が加工して作成"
TIDY = DATA / "pop_age_sex.csv"
ERA = {"大正": 1911, "昭和": 1925, "平成": 1988, "令和": 2018}


def _rows(path):
    b = Path(path).read_bytes()
    if b[:2] == b"PK":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
        return {n: [list(r) for r in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames}
    import xlrd
    w = xlrd.open_workbook(file_contents=b)
    return {n: [w.sheet_by_name(n).row_values(i) for i in range(w.sheet_by_name(n).nrows)] for n in w.sheet_names()}


def _year(s):
    s = re.sub(r"\s", "", str(s))
    m = re.match(r"(大正|昭和|平成|令和)(元|\d+)年", s)
    if m:
        return ERA[m.group(1)] + (1 if m.group(2) == "元" else int(m.group(2)))
    m = re.match(r"^(19|20)\d\d", s)
    return int(m.group(0)) if m else None


def _age(s):
    s = re.sub(r"\s", "", str(s))
    if s in ("", "None"):
        return None
    if s.startswith("総数"):
        return "total"
    m = re.match(r"^(\d+)(歳以上)$", s)
    if m:
        return ("top", int(m.group(1)))
    m = re.match(r"^(\d+)(\.0)?(歳)?$", s)
    return int(m.group(1)) if m else None


def _num(v):
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return np.nan


def parse_longterm(path, want_sheet=lambda n: True):
    out = []
    for name, R in _rows(path).items():
        if not want_sheet(name):
            continue
        hdr = None
        for ri in (5, 6):
            ys = {c: _year(v) for c, v in enumerate(R[ri]) if _year(v)}
            if len(ys) >= 3:
                hdr = ys; break
        if not hdr:
            continue
        # 年齢の列：表頭に「齢」「Age」「数え年」がある列（なければ0列目）
        age_cols = sorted({c for ri in range(6, 11) for c, v in enumerate(R[ri]) if any(k in str(v) for k in ("齢", "Age", "数 え 年", "Calendar"))} | {0})
        kazoe = {c for ri in range(6, 11) for c, v in enumerate(R[ri]) if "数 え 年" in str(v) or "Calendar" in str(v)}
        for c, y in hdr.items():
            ac = max(a for a in age_cols if a < c)
            for r in R[11:]:
                if ac >= len(r) or c + 2 >= len(r):
                    continue
                a = _age(r[ac])
                if a is None:
                    continue
                vals = [_num(r[c + k]) for k in range(3)]
                if all(np.isnan(vals)):
                    continue
                if a == "total":
                    age, top = -1, 0
                elif isinstance(a, tuple):
                    age, top = a[1], 1
                else:
                    age, top = a, 0
                out.append(dict(year=y, age=age, top=top, kazoe=int(ac in kazoe), both=vals[0], male=vals[1], female=vals[2], sheet=name))
    return pd.DataFrame(out)


def parse_annual(path):
    R = list(_rows(path).values())[0]
    out = []
    for r in R:
        k = next((i for i, v in enumerate(r) if "10月1日" in str(v)), None)
        if k is None:
            continue
        y = int(str(r[k])[:4]); a = _age(r[k + 1])
        if a is None:
            continue
        age, top = (-1, 0) if a == "total" else ((a[1], 1) if isinstance(a, tuple) else (a, 0))
        out.append(dict(year=y, age=age, top=top, kazoe=0, both=_num(r[k + 2]), male=_num(r[k + 3]), female=_num(r[k + 4]), sheet="各年"))
    return pd.DataFrame(out)


def build():
    a = parse_longterm(RAW / "140_1920_2000_kakusai.k0")
    b = parse_longterm(RAW / "140_2000_2020_kakusai.k4", want_sheet=lambda n: n.startswith("総人口"))
    c = pd.concat([parse_annual(p) for p in sorted(RAW.glob("140_20[2-9][0-9].k0"))])
    df = pd.concat([a[a.year < 2000], b, c[c.year > 2020]])
    df = df.drop_duplicates(["year", "age", "top"], keep="last").sort_values(["year", "age"])
    df = df[["year", "age", "top", "kazoe", "both", "male", "female"]]
    df.to_csv(TIDY, index=False)
    return df


def load():
    if TIDY.exists():
        return pd.read_csv(TIDY)
    return build()


def topcode(df, cap=85):
    """年齢をcap歳以上でまとめる（年によって最上位の区分が違うため）"""
    d = df[df.age >= 0].copy()
    d["age"] = d.age.clip(upper=cap)
    return d.groupby(["year", "age"], as_index=False)[["both", "male", "female"]].sum()
