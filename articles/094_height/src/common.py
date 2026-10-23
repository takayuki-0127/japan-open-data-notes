"""共通：文部科学省「学校保健統計調査」（e-Stat）の年次統計・都道府県表を読む"""
from pathlib import Path
import re, logging, warnings
import numpy as np, pandas as pd
import matplotlib.pyplot as plt
from matplotlib import font_manager
warnings.filterwarnings("ignore")
logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
for f in font_manager.findSystemFonts():
    if "NotoSansCJK" in f or "NotoSans-CJK" in f:
        font_manager.fontManager.addfont(f)
plt.rcParams.update({"font.family": ["Noto Sans CJK JP", "IPAexGothic", "Yu Gothic", "Meiryo", "sans-serif"],
                     "axes.spines.top": False, "axes.spines.right": False, "font.size": 11, "axes.titlesize": 13,
                     "figure.dpi": 150, "savefig.bbox": "tight"})
ROOT = Path(__file__).resolve().parents[1]
RAW, FIG, RES = ROOT / "data" / "raw", ROOT / "figures", ROOT / "results"
if not (RAW / "gakko_nenji_0.xlsx").exists() and (ROOT / "data" / "gakko_raw.zip").exists():
    import zipfile
    RAW.mkdir(parents=True, exist_ok=True)
    zipfile.ZipFile(ROOT / "data" / "gakko_raw.zip").extractall(RAW)
FIG.mkdir(exist_ok=True); RES.mkdir(exist_ok=True)
BLUE, ORANGE, GRAY, INK, LIGHT, GREEN, PURPLE = "#2a78d6", "#eb6834", "#a3a29c", "#0b0b0b", "#a9c9ef", "#3aa07a", "#7a5bd6"
SOURCE = "出典：文部科学省「学校保健統計調査」（e-Stat、https://www.e-stat.go.jp/stat-search/files?toukei=00400002）を加工して作成"
ERA = {"明治": 1867, "大正": 1911, "昭和": 1925, "平成": 1988, "令和": 2018}
COVID = (2020, 2023)  # 令和2〜5年度は測定時期が違うため、過去と直接比べられない


def _num(v):
    return float(v) if isinstance(v, (int, float)) else np.nan


def _rows(path, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def era_years(labels):
    """「明治33年度」「　　34年」「令和元年」… を西暦の年度に直す"""
    out, era = [], None
    for lab in labels:
        s = str(lab or "").replace("　", " ").strip()
        m = re.match(r"^(明治|大正|昭和|平成|令和)?\s*(元|\d+)\s*年", s)
        if not m:
            out.append(None); continue
        if m.group(1):
            era = m.group(1)
        n = 1 if m.group(2) == "元" else int(m.group(2))
        out.append(ERA[era] + n if era else None)
    return out


def load_height():
    """年齢別 平均身長（年度×性×年齢、cm）"""
    recs = []
    for sh, sex in [(0, "男"), (1, "女")]:
        rows = _rows(RAW / "gakko_nenji_0.xlsx", sh)
        ages = [int(str(a).replace("歳", "")) if a and "歳" in str(a) else None for a in rows[4]]
        years = era_years([r[0] for r in rows])
        for r, y in zip(rows, years):
            if y is None:
                continue
            for j, a in enumerate(ages):
                if a is not None:
                    recs.append((y, sex, a, _num(r[j])))
    return pd.DataFrame(recs, columns=["year", "sex", "age", "height"]).dropna()


def _vcols(h4, h5):
    """列名（2段目・3段目）から使う列の位置を探す"""
    c = {"受検者計": 1, "1.0未満": 10, "0.7以上1.0未満": 11, "0.3以上0.7未満": 12, "0.3未満": 13}
    for j, (a, b) in enumerate(zip(h4, h5)):
        if a == "むし歯（う歯）" and b == "計": c["むし歯"] = j
        if b == "未処置歯のある者": c["むし歯未処置"] = j
        if b in ("眼の疾病・異常", "鼻・副鼻腔疾患", "アトピー性皮膚炎", "ぜん息", "栄養状態"): c[b] = j
    assert h5[10] == "計" and h5[13] == "0.3未満" and "1.0未満" not in (h5[10],)
    return c


def load_disease():
    """学校種別 疾病・異常被患率等の推移（年度×学校種×性、％）"""
    import openpyxl
    recs = []
    wb = openpyxl.load_workbook(RAW / "gakko_nenji_5.xlsx", read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        t = str(rows[2][0])
        school = re.search(r"(幼稚園|小学校|中学校|高等学校)", t).group(1)
        sex = re.search(r"（\d）(計|男|女)", t).group(1)
        cols = _vcols(rows[4], rows[5])
        for r, y in zip(rows, era_years([r[0] for r in rows])):
            if y is None:
                continue
            d = {"year": y, "school": school, "sex": sex}
            for name, j in cols.items():
                d[name] = _num(r[j])
            recs.append(d)
    return pd.DataFrame(recs)


def load_pref(year):
    """都道府県別 年齢別 疾病・異常被患率等（年齢ごと）：裸眼視力1.0未満・0.3未満・むし歯
    元の Excel が data/raw に無ければ、整えた data/pref_vision.csv を使う"""
    if not (RAW / f"gakko_pref{year}_3.xlsx").exists():
        p = pd.read_csv(ROOT / "data" / "pref_vision.csv")
        return p[p.year == year].reset_index(drop=True)
    import openpyxl
    wb = openpyxl.load_workbook(RAW / f"gakko_pref{year}_3.xlsx", read_only=True, data_only=True)
    recs = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        m = re.search(r"(\d+)歳\s*（\d）\s*(計|男|女)", str(rows[1][0]).replace("　", " "))
        age, sex = int(m.group(1)), m.group(2)
        jm = [j for j, (a, b) in enumerate(zip(rows[4], rows[5])) if a == "むし歯（う歯）" and b == "計"]
        for r in rows[6:]:
            if not r[0] or not isinstance(r[1], (int, float)):
                continue
            recs.append({"year": year, "pref": str(r[0]).strip(), "age": age, "sex": sex,
                         "1.0未満": _num(r[10]), "0.3未満": _num(r[13]), "むし歯": _num(r[jm[0]])})
    return pd.DataFrame(recs)


def shade_covid(ax):
    ax.axvspan(COVID[0] - .5, COVID[1] + .5, color=GRAY, alpha=.18, lw=0)


def note(fig, extra=""):
    fig.text(0.01, -0.02, (extra + "\n" if extra else "") + SOURCE, fontsize=7.5, color="#555", ha="left", va="top")
