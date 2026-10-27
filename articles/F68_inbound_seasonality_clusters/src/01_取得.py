"""01 取得：JNTO「国籍/月別 訪日外客数（2003年～2026年）」の Excel を、国・地域×年×月の縦長の表にする。

このリポジトリでは、ブラウザの中で SheetJS を使って同じ処理を行い data/raw/jnto_monthly_2003_2026.csv を作った
（コンテナから JNTO のサイトに直接つながらなかったため）。手元に Excel を置けば、このスクリプトで同じ表を作れる。

  python src/01_取得.py path/to/20261021_1615-5.xlsx

Excel の掲載ページ：https://www.jnto.go.jp/statistics/data/visitors-statistics/
"""
import re
import sys

import pandas as pd
from openpyxl import load_workbook

from common import RAW

def main(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows_out = []
    for sn in wb.sheetnames:
        rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        hi = next(i for i, r in enumerate(rows) if r and "1月" in r)
        mc = {int(re.match(r"(\d+)月", str(c)).group(1)): j for j, c in enumerate(rows[hi]) if c and re.fullmatch(r"\d+月", str(c))}
        first = min(mc.values())
        parent = ""
        for r in rows[hi + 1:]:
            if not r:
                continue
            lab = [str(x).strip() for x in r[:first] if x not in (None, "")]
            if not lab:
                continue
            name = lab[-1]
            if name.startswith("注"):
                break
            if r[0] not in (None, ""):
                parent = str(r[0]).strip()
            for m, j in mc.items():
                v = r[j] if j < len(r) else None
                if isinstance(v, (int, float)):
                    rows_out.append((int(sn), name, parent, m, v))
    d = pd.DataFrame(rows_out, columns=["year", "market", "parent", "month", "visitors"])
    d.to_csv(RAW / "jnto_monthly_2003_2026.csv", index=False)
    print(len(d), "rows")

if __name__ == "__main__":
    main(sys.argv[1])
