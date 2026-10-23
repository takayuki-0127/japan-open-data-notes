"""共通：厚生労働省「一般職業紹介状況（職業安定業務統計）」の長期時系列表（e-Stat）を読む"""
from pathlib import Path
import re, logging, warnings
import numpy as np, pandas as pd
import matplotlib.pyplot as plt
from matplotlib import font_manager
warnings.filterwarnings("ignore")
logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
for f in font_manager.findSystemFonts():
    if "NotoSansCJK" in f or "NotoSans-CJK" in f:
        font_manager.fontManager.addfont(f)
plt.rcParams.update({"font.family": ["Noto Sans CJK JP", "IPAexGothic", "Yu Gothic", "Meiryo", "sans-serif"],
                     "axes.spines.top": False, "axes.spines.right": False, "font.size": 11, "axes.titlesize": 13,
                     "figure.dpi": 150, "savefig.bbox": "tight"})
ROOT = Path(__file__).resolve().parents[1]
DATA, RAW, FIG, RES = ROOT / "data", ROOT / "data" / "raw", ROOT / "figures", ROOT / "results"
FIG.mkdir(exist_ok=True); RES.mkdir(exist_ok=True)
BLUE, ORANGE, GRAY, INK, LIGHT, GREEN, PURPLE = "#2a78d6", "#eb6834", "#a3a29c", "#0b0b0b", "#a9c9ef", "#3aa07a", "#7a5bd6"
SOURCE = "出典：厚生労働省「一般職業紹介状況（職業安定業務統計）」（e-Stat）を加工して作成"
PREFS = ["北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県", "茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県",
         "新潟県", "富山県", "石川県", "福井県", "山梨県", "長野県", "岐阜県", "静岡県", "愛知県", "三重県", "滋賀県", "京都府", "大阪府", "兵庫県",
         "奈良県", "和歌山県", "鳥取県", "島根県", "岡山県", "広島県", "山口県", "徳島県", "香川県", "愛媛県", "高知県", "福岡県", "佐賀県", "長崎県",
         "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県"]
# 元の Excel（e-Stat の長期時系列表）。data/raw/ に置くと data/*.csv を作り直せる
RAW_FILES = {"occ_new": "kyujin_21.xlsx",       # 第21表 職業別（平成21年改定 日本標準職業分類）（令和4年4月～）
             "occ_old": "kyujin_old400_01.xlsx",  # 旧第21表 職業別（平成23年改定 厚生労働省編職業分類）（平成24年3月～令和5年3月）
             "pref_jyuri": "kyujin_11.xlsx",     # 第11表 都道府県（受理地）別
             "pref_shugyo": "kyujin_18.xlsx",    # 第18表 都道府県（就業地）別
             "pref_hw": "kyujin_26.xlsx"}        # 第26表 都道府県（受理地）別（ハローワーク利用登録者）
OCC_MEAS = {1: "新規求人", 2: "有効求人", 3: "新規求職", 4: "有効求職", 5: "就職件数"}  # 第21表-1〜5（パートを含む常用）
MON = {"１月": 1, "２月": 2, "３月": 3, "４月": 4, "５月": 5, "６月": 6, "７月": 7, "８月": 8, "９月": 9, "10月": 10, "11月": 11, "12月": 12}


def _z2h(s):
    return str(s).translate(str.maketrans("０１２３４５６７８９", "0123456789")).strip()


def _rows(path, idx):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return [list(r) for r in wb[wb.sheetnames[idx]].iter_rows(values_only=True)]


def _occ_sheet(path, idx, tag):
    """職業別の表（行＝職業、列＝年・年度・月）から月の値だけを縦長にする"""
    rows = _rows(path, idx)
    yrs, mons = rows[1], rows[3]
    out = []
    for i, r in enumerate(rows[5:5 + 69]):
        name = str(r[0]).strip()
        if name.startswith("（注"):
            break
        for j in range(1, len(r)):
            m = _z2h(mons[j]) if mons[j] else ""
            if not m.endswith("月") or r[j] in (None, "", "-", "－"):
                continue
            y = int(re.match(r"(\d{4})", str(yrs[j])).group(1))
            out.append((tag, i, name, f"{y}-{int(m[:-1]):02d}", float(r[j])))
    return out


def build_occ():
    recs = []
    for tag, key in (("旧", "occ_old"), ("新", "occ_new")):
        for idx, meas in OCC_MEAS.items():
            for t in _occ_sheet(RAW / RAW_FILES[key], idx - 1, tag):
                recs.append(t + (meas,))
    d = pd.DataFrame(recs, columns=["分類", "行", "職業", "年月", "値", "指標"])
    names = d[["分類", "行", "職業"]].drop_duplicates().sort_values(["分類", "行"])
    w = d.pivot_table(index=["分類", "行", "年月"], columns="指標", values="値").reset_index()
    w = w[["分類", "行", "年月"] + list(OCC_MEAS.values())]
    for c in OCC_MEAS.values():
        w[c] = w[c].round().astype("Int64")
    names.to_csv(DATA / "occ_names.csv", index=False)
    w.to_csv(DATA / "occ_monthly.csv", index=False)


def _pref_sheet(path, idx, with_nation=True):
    rows = _rows(path, idx)
    head = [str(c).strip() if c else "" for c in rows[1]]
    cols = [(j, h) for j, h in enumerate(head) if h in PREFS or h == "全国"]
    ann, mon = [], []
    for r in rows[3:]:
        if not r[0]:
            continue
        lab = str(r[0]).strip()
        m = re.match(r"(\d{4})年(度?)$", lab)
        if not m:
            continue
        y = int(m.group(1))
        for j, h in cols:
            v = r[j]
            if v in (None, "", "-", "－", "…"):
                continue
            try:
                v = float(v)
            except ValueError:
                continue
            if r[2] in (None, ""):
                if not m.group(2):
                    ann.append((h, y, v))
            else:
                mm = MON.get(str(r[2]).strip())
                if mm:
                    mon.append((h, f"{y}-{mm:02d}", v))
    return ann, mon


def build_pref():
    # 受理地：第11表-3 有効求人数、-4 有効求職者数、-6 有効求人倍率（実数）
    A = {}
    M = {}
    for idx, k in ((2, "有効求人"), (3, "有効求職"), (5, "有効求人倍率")):
        a, m = _pref_sheet(RAW / RAW_FILES["pref_jyuri"], idx)
        A[k] = pd.DataFrame(a, columns=["都道府県", "年", k]).set_index(["都道府県", "年"])
        M[k] = pd.DataFrame(m, columns=["都道府県", "年月", k]).set_index(["都道府県", "年月"])
    ann = pd.concat(A.values(), axis=1).reset_index().assign(基準="受理地")
    # 就業地：第18表-2 有効求人数、-4 有効求人倍率（実数）。有効求職者数は受理地と同じもの
    B = {}
    for idx, k in ((1, "有効求人"), (3, "有効求人倍率")):
        a, m = _pref_sheet(RAW / RAW_FILES["pref_shugyo"], idx)
        B[k] = pd.DataFrame(a, columns=["都道府県", "年", k]).set_index(["都道府県", "年"])
    sh = pd.concat(B.values(), axis=1).reset_index().assign(基準="就業地")
    sh = sh.merge(ann[["都道府県", "年", "有効求職"]], on=["都道府県", "年"], how="left")
    pa = pd.concat([ann, sh], ignore_index=True)[["基準", "都道府県", "年", "有効求人", "有効求職", "有効求人倍率"]]
    pa.to_csv(DATA / "pref_annual.csv", index=False)
    # 月次（2019年1月〜）：受理地の有効求人・有効求職と、ハローワーク利用登録者のみの有効求職（第26表-2、2021年9月〜）
    mon = pd.concat(M.values(), axis=1).reset_index()
    _, hw = _pref_sheet(RAW / RAW_FILES["pref_hw"], 1)
    hw = pd.DataFrame(hw, columns=["都道府県", "年月", "有効求職_HW利用登録者"])
    mon = mon.merge(hw, on=["都道府県", "年月"], how="left")
    mon = mon[mon["年月"] >= "2019-01"].sort_values(["都道府県", "年月"])
    mon.to_csv(DATA / "pref_monthly.csv", index=False)


def load_occ():
    if not (DATA / "occ_monthly.csv").exists():
        build_occ()
    return pd.read_csv(DATA / "occ_monthly.csv"), pd.read_csv(DATA / "occ_names.csv")


def load_pref():
    if not (DATA / "pref_annual.csv").exists():
        build_pref()
    return pd.read_csv(DATA / "pref_annual.csv"), pd.read_csv(DATA / "pref_monthly.csv")


def foot(fig, note=""):
    fig.text(0.0, -0.02, (note + "\n" if note else "") + SOURCE, fontsize=8, color=GRAY, ha="left", va="top")


if __name__ == "__main__":
    build_occ(); build_pref()
