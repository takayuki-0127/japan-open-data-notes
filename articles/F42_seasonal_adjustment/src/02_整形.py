"""02 整形：原数値と公表の季節調整値を1つの表にする
出力: data/series.csv（date, series, nsa, sa_pub）、data/lfs_factors_pub.csv（労働力調査が公表している季節指数）
"""
import openpyxl
import pandas as pd
from common import RAW, DATA

LFS_COLS = {"労働力人口": 4, "就業者": 7, "雇用者": 10, "完全失業者": 13, "非労働力人口": 16, "完全失業率": 19}


def _num(v):
    if isinstance(v, str):
        v = v.strip("<>() ")
        try:
            return float(v)
        except ValueError:
            return None
    return float(v) if isinstance(v, (int, float)) else None


def read_lfs(sheet):
    ws = openpyxl.load_workbook(RAW / "lfs.xlsx", data_only=True, read_only=True)[sheet]
    rows, year = [], None
    for r in ws.iter_rows(min_row=11, values_only=True):
        if isinstance(r[0], int):
            year = r[0]
        elif isinstance(r[0], str) and r[0][:2] in ("昭和", "平成", "令和"):
            era = {"昭和": 1925, "平成": 1988, "令和": 2018}[r[0][:2]]
            year = era + int("".join(ch for ch in r[0][2:] if ch.isdigit()) or 1)
        if r[1] and str(r[1]).replace("*", "").strip().endswith("月") and year:
            m = int(str(r[1]).replace("*", "").replace("月", "").strip())
            rec = {"date": pd.Timestamp(year, m, 1)}
            for k, c in LFS_COLS.items():
                rec[k] = _num(r[c])
            rows.append(rec)
            if m == 12:
                year += 1
    return pd.DataFrame(rows).set_index("date")


def read_jobs():
    ws = openpyxl.load_workbook(RAW / "jobs.xlsx", data_only=True, read_only=True).worksheets[0]
    sa, nsa = {}, {}
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[0] and str(r[0]).endswith("年") and str(r[0])[:4].isdigit():
            y = int(str(r[0])[:4])
            for m in range(1, 13):
                if isinstance(r[1 + m], (int, float)):
                    nsa[pd.Timestamp(y, m, 1)] = float(r[1 + m])
                if isinstance(r[19 + m], (int, float)):
                    sa[pd.Timestamp(y, m, 1)] = float(r[19 + m])
    return pd.Series(nsa), pd.Series(sa)


def read_iip():
    d = pd.read_csv(RAW / "iip_mining_manufacturing.csv", dtype={"ym": str})
    d["date"] = pd.to_datetime(d["ym"], format="%Y%m")
    d = d.set_index("date")
    out = {}
    for p, nm in [("prod", "鉱工業生産"), ("ship", "鉱工業出荷")]:
        nsa = d[f"{p}_nsa"].combine_first(d[f"{p}_nsa_c"])   # 2018年以降は2020年基準、それ以前は接続指数
        sa = d[f"{p}_sa"].combine_first(d[f"{p}_sa_c"])
        out[nm] = (nsa, sa)
    return out


def main():
    rows = []
    sa, nsa = read_lfs("季節調整値"), read_lfs("原数値")
    for k in LFS_COLS:
        df = pd.DataFrame({"nsa": nsa[k], "sa_pub": sa[k]}).dropna(subset=["nsa"])
        df = df[df.index >= "1990-01-01"]
        df["series"] = k
        rows.append(df)
    jn, js = read_jobs()
    df = pd.DataFrame({"nsa": jn, "sa_pub": js})
    df = df[df.index >= "1990-01-01"]
    df["series"] = "有効求人倍率"
    rows.append(df)
    for k, (n, s) in read_iip().items():
        df = pd.DataFrame({"nsa": n, "sa_pub": s})
        df["series"] = k
        rows.append(df)
    out = pd.concat(rows).reset_index().rename(columns={"index": "date"})
    out = out[["date", "series", "nsa", "sa_pub"]].dropna(subset=["nsa"])
    out.to_csv(DATA / "series.csv", index=False)
    print(out.groupby("series")["date"].agg(["min", "max", "count"]))

    f = read_lfs("季節指数").dropna(how="all")
    f.to_csv(DATA / "lfs_factors_pub.csv")
    print("季節指数（公表）:", f.index.min().date(), "〜", f.index.max().date(), len(f))


if __name__ == "__main__":
    main()
