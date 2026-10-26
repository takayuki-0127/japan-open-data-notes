"""02 整形：月次の表 data/monthly.csv を作る"""
import json
import openpyxl
import pandas as pd
from common import RAW, DATA


def read_lfs(sheet):
    ws = openpyxl.load_workbook(RAW / "lfs.xlsx", data_only=True, read_only=True)[sheet]
    out, year = [], None
    for r in ws.iter_rows(min_row=11, values_only=True):
        if isinstance(r[0], int):
            year = r[0]
        elif isinstance(r[0], str) and r[0].startswith("昭和28"):
            year = 1953
        if r[1] and str(r[1]).replace("*", "").strip().endswith("月") and year:
            m = int(str(r[1]).replace("*", "").replace("月", "").strip())
            v = r[19]
            if isinstance(v, str):
                v = v.strip("<>() ")
                try:
                    v = float(v)
                except ValueError:
                    v = None
            if isinstance(v, (int, float)):
                out.append((pd.Timestamp(year, m, 1), float(v)))
            if m == 12:
                year += 1
    return pd.Series(dict(out))


def read_jobs():
    ws = openpyxl.load_workbook(RAW / "jobs.xlsx", data_only=True, read_only=True).worksheets[0]
    sa, nsa = {}, {}
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[0] and str(r[0]).endswith("年") and str(r[0])[:4].isdigit():
            y = int(str(r[0])[:4])
            for m in range(1, 13):
                if isinstance(r[1 + m], (int, float)):
                    nsa[pd.Timestamp(y, m, 1)] = float(r[1 + m])
                if isinstance(r[19 + m], (int, float)):
                    sa[pd.Timestamp(y, m, 1)] = float(r[19 + m])
    return pd.Series(sa), pd.Series(nsa)


def read_ci():
    d = json.loads((RAW / "f13_ci.json").read_text(encoding="utf-8"))

    def table(rows, cols):
        out = {}
        for r in rows:
            if len(r) > 3 and isinstance(r[1], (int, float)) and isinstance(r[2], (int, float)):
                t = pd.Timestamp(int(r[1]), int(r[2]), 1)
                out[t] = {k: (float(r[i]) if i < len(r) and isinstance(r[i], (int, float)) else None) for k, i in cols.items()}
        return pd.DataFrame(out).T

    ci = table(d["1026ci.xlsx"], {"ci_lead": 3, "ci_coin": 4})
    c2 = table(d["1026ci2.xlsx"], {"iip": 3, "labor_input": 6})
    c1 = table(d["1026ci1.xlsx"], {"new_offers": 5, "consumer_conf": 8})
    return ci.join(c2, how="outer").join(c1, how="outer")


if __name__ == "__main__":
    u_sa, u_nsa = read_lfs("季節調整値"), read_lfs("原数値")
    j_sa, j_nsa = read_jobs()
    df = pd.DataFrame({"u_sa": u_sa, "u_nsa": u_nsa, "jobs_sa": j_sa, "jobs_nsa": j_nsa})
    df = df.join(read_ci(), how="outer").sort_index()
    df = df[df.index >= "1975-01-01"]
    df.index.name = "month"
    df.to_csv(DATA / "monthly.csv", float_format="%.4g")
    print(df.tail(3).T)
    print(df.notna().sum())
    print("first full row:", df.dropna().index.min(), "last u:", df["u_sa"].dropna().index.max())
